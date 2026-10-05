"""Independent read-only EIA/DOE source-cell and factual-projection QA."""
from collections import Counter, defaultdict
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

import openpyxl

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / ".runtime"))
import xlrd

ENERGY = BASE / "energy_sources"
OUT = ENERGY / "normalized"
REPORT = Path(__file__).resolve().parent / "energy_validation.json"


def sha(path):
    result = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def read_jsonl(path):
    with Path(path).open(encoding="utf-8") as stream:
        return [json.loads(line) for line in stream if line.strip()]


def clean(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return value.strip() if isinstance(value, str) else value


def source_key(source):
    return source["archive"], source["member"], source["sheet"], source["row"]


def identifier(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def doe_date(cell, workbook):
    if cell.ctype == xlrd.XL_CELL_DATE:
        year, month, day, *_ = xlrd.xldate_as_tuple(cell.value, workbook.datemode)
        return dt.date(year, month, day).isoformat()
    text = str(cell.value).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def main():
    summary = json.loads((OUT / "collection_summary.json").read_text(encoding="utf-8"))
    if "incidents_with_unknown_start_date" not in summary:
        raise RuntimeError("Wait for completed corrected normalizer with explicit unknown-start-date count")
    errors, examples = Counter(), defaultdict(list)
    def check(condition, code, detail):
        if not condition:
            errors[code] += 1
            if len(examples[code]) < 8:
                examples[code].append(detail)
    observations = read_jsonl(OUT / "annual_observations.jsonl")
    projected = {source_key(record["source"]): record for record in observations}
    check(len(projected) == len(observations), "duplicate_source_row_projection", "annual source locators must be unique")
    source_counts, table_boundaries, material_sha, source_rows = {}, [], {}, {}
    skipped_nonblank = []
    verified_cells = 0
    prefixes = {"utility": "1___Utility", "plant": "2___Plant", "generator": "3_1_Generator", "ownership": "4___Owner"}
    for year in range(2018, 2026):
        archive = ENERGY / f"raw/eia860{year}.zip"
        archive_hash = sha(archive)
        material_sha[str(archive.relative_to(ENERGY))] = archive_hash
        counts = Counter()
        with zipfile.ZipFile(archive) as z:
            check(z.testzip() is None, "archive_crc", str(archive))
            for kind, prefix in prefixes.items():
                members = [name for name in z.namelist() if name.startswith(prefix) and name.endswith(".xlsx")]
                check(len(members) == 1, "expected_official_table_member", f"{year}:{kind}:{members}")
                for member in members:
                    workbook = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True, data_only=True)
                    for sheet in workbook:
                        header, header_row, kept = None, None, []
                        for row_num, values in enumerate(sheet.iter_rows(values_only=True), 1):
                            if header is None:
                                if "Utility ID" in values:
                                    header = [str(v).strip() if v is not None else "" for v in values]
                                    header_row = row_num
                                continue
                            cells = {key: clean(value) for key, value in zip(header, values) if key}
                            uid = identifier(cells.get("Utility ID"))
                            key = f"raw/eia860{year}.zip", member, sheet.title, row_num
                            if uid is None:
                                if any(value is not None and value != "" for value in values):
                                    skipped_nonblank.append({"source": key, "has_asset_or_owner_cells": any(cells.get(field) not in (None, "") for field in ("Plant Code", "Generator ID", "Owner Name"))})
                                continue
                            counts[kind] += 1
                            source_rows[key] = {"kind": kind, "year": year, "cells": cells}
                            kept.append(key)
                            record = projected.get(key)
                            check(record is not None, "missing_eia_data_row", list(key))
                            if record is None:
                                continue
                            check(record["survey_year"] == year and record["table_kind"] == kind,
                                  "eia_projection_identity", list(key))
                            check(record["source"].get("archive_sha256") == archive_hash, "eia_projection_source_hash", list(key))
                            for field, value in record["fields"].items():
                                verified_cells += 1
                                check(field in cells and value == cells.get(field), "eia_projected_cell_changed", {"source": list(key), "field": field, "normalized": value, "raw": cells.get(field)})
                        check(header is not None, "missing_eia_header", f"{year}:{member}:{sheet.title}")
                        if kept:
                            table_boundaries.append({"year": year, "kind": kind, "member": member, "sheet": sheet.title,
                                "header_row": header_row, "source_rows": len(kept), "first_source_row": kept[0][-1],
                                "last_source_row": kept[-1][-1], "first_projection_present": kept[0] in projected,
                                "last_projection_present": kept[-1] in projected})
                    workbook.close()
        source_counts[str(year)] = dict(counts)
        check(dict(counts) == summary["annual_table_rows"].get(str(year)), "eia_annual_row_count", {"year": year, "actual": dict(counts), "summary": summary["annual_table_rows"].get(str(year))})
    check(set(projected) == set(source_rows), "extra_or_missing_eia_projected_locators", {"extra": len(set(projected) - set(source_rows)), "missing": len(set(source_rows) - set(projected))})
    check(not any(record["has_asset_or_owner_cells"] for record in skipped_nonblank), "eia_nonblank_data_without_utility_id", skipped_nonblank[:8])

    entities = read_jsonl(OUT / "entities.jsonl"); relations = read_jsonl(OUT / "relations.jsonl")
    entity_map = {record["entity_id"]: record for record in entities}
    check(len(entity_map) == len(entities), "duplicate_entity_ids", "IDs must be unique")
    allowed_types = {"utility_reporting_entity", "plant", "generator", "region_state_code", "industry_naics_code", "fuel_code"}
    check({record["entity_type"] for record in entities} <= allowed_types, "unsupported_or_paper_type_relabelling", "Assets must not be relabelled as Products or ultimate parent Firms")
    expected_entities, expected_relations = {}, {}
    ownership_source_cells = defaultdict(list)
    def observed_entity(code, typ, year):
        if code:
            expected_entities.setdefault(code, {"type": typ, "years": set()})["years"].add(year)
    def observed_relation(head, tail, kind, year):
        if head and tail:
            expected_relations.setdefault((head, tail, kind), set()).add(year)
    for item in source_rows.values():
        fields, year, kind = item["cells"], item["year"], item["kind"]
        uid, pid, gid = map(identifier, (fields.get("Utility ID"), fields.get("Plant Code"), fields.get("Generator ID")))
        utility = "eia:utility:" + uid if uid else None
        plant = "eia:plant:" + pid if pid else None
        generator = f"eia:generator:{pid}:{gid}" if pid and gid else None
        observed_entity(utility, "utility_reporting_entity", year)
        observed_entity(plant, "plant", year)
        observed_entity(generator, "generator", year)
        if kind == "plant":
            observed_relation(plant, utility, "ReportedBy", year)
        if generator:
            observed_relation(generator, plant, "LocatedAtPlant", year)
            if kind == "generator":
                observed_relation(generator, utility, "ReportedBy", year)
        state = identifier(fields.get("State"))
        if state:
            region = "eia:state:" + state
            observed_entity(region, "region_state_code", year)
            observed_relation(plant if plant else utility, region, "LocatedInState", year)
        if kind == "plant":
            naics = identifier(fields.get("Primary Purpose (NAICS Code)"))
            if naics:
                industry = "eia:naics:" + naics
                observed_entity(industry, "industry_naics_code", year)
                observed_relation(plant, industry, "PrimaryPurposeNAICS", year)
        if kind == "generator":
            for field in (f"Energy Source {i}" for i in range(1, 7)):
                fuel = identifier(fields.get(field))
                if fuel:
                    target = "eia:fuel:" + fuel
                    observed_entity(target, "fuel_code", year)
                    observed_relation(generator, target, "UsesEnergySource", year)
        if kind == "ownership" and generator:
            owner_id = identifier(fields.get("Ownership ID"))
            if owner_id:
                owner = "eia:utility:" + owner_id
                observed_entity(owner, "utility_reporting_entity", year)
                observed_relation(owner, generator, "OwnsReportedGenerator", year)
                ownership_source_cells[(owner, generator, year)].append(fields.get("Percent Owned"))
    check(set(entity_map) == set(expected_entities), "extra_or_missing_source_supported_entities", {"extra": len(set(entity_map) - set(expected_entities)), "missing": len(set(expected_entities) - set(entity_map))})
    for code, record in entity_map.items():
        expected = expected_entities.get(code)
        if expected:
            check(record["entity_type"] == expected["type"] and record.get("is_synthetic") is False,
                  "entity_kind_or_provenance", code)
            check(record["first_observed_survey_year"] == min(expected["years"]) and record["last_observed_survey_year"] == max(expected["years"]), "entity_year_coverage", code)
        check(source_key(record["first_source"]) in source_rows, "entity_first_source_missing", code)
    relation_map = {(record["head"], record["tail"], record["relation_type"]): record for record in relations}
    check(len(relation_map) == len(relations), "duplicate_relation_triples", "triples must be unique")
    check(set(relation_map) == set(expected_relations), "extra_or_missing_source_supported_relations", {"extra": len(set(relation_map) - set(expected_relations)), "missing": len(set(expected_relations) - set(relation_map))})
    for key, record in relation_map.items():
        check(key[0] in entity_map and key[1] in entity_map, "relation_endpoint_missing", key)
        check(record.get("is_synthetic") is False, "relation_synthetic_flag", key)
        if key in expected_relations:
            years = expected_relations[key]
            check(record["first_observed_survey_year"] == min(years) and record["last_observed_survey_year"] == max(years), "relation_year_coverage", key)
        check(source_key(record["first_source"]) in source_rows, "relation_first_source_missing", key)
        for year, attributes in record.get("observed_attributes", {}).items():
            matches = ownership_source_cells[(key[0], key[1], int(year))]
            check(attributes.get("source_percent_owned_cell") in matches, "ownership_cell_changed", key)

    incidents = read_jsonl(OUT / "risk_events.jsonl")
    incident_map = {(record["source"]["file"], record["source"]["sheet"], record["source"]["row"]): record for record in incidents}
    check(len(incidents) == len(incident_map), "duplicate_doe_source_rows", "incidents must have unique source locators")
    doe_counts, missing_dates, unknown_numeric_values, raw_incident_keys = {}, [], Counter(), set()
    for path in sorted((ENERGY / "raw").glob("doe417_*_Annual_Summary.xls")):
        year = int(path.name.split("_")[1]); raw_count = 0
        digest = sha(path); material_sha[str(path.relative_to(ENERGY))] = digest
        book = xlrd.open_workbook(path, logfile=io.StringIO())
        for sheet in book.sheets():
            header_row = next((i for i in range(min(15, sheet.nrows)) if "Date Event Began" in sheet.row_values(i)), None)
            if header_row is None:
                continue
            headers = [str(value).strip() for value in sheet.row_values(header_row)]
            col = headers.index("Date Event Began")
            boundaries = []
            for i in range(header_row + 1, sheet.nrows):
                raw = {name: clean(value) for name, value in zip(headers, sheet.row_values(i)) if name}
                if not any(raw.get(field) not in (None, "") for field in ("Area Affected", "Event Type", "Alert Criteria")):
                    continue
                raw_count += 1
                key = path.relative_to(ENERGY).as_posix(), sheet.name, i + 1
                raw_incident_keys.add(key); boundaries.append(i + 1)
                record = incident_map.get(key)
                check(record is not None, "missing_doe_incident_row", list(key))
                began = doe_date(sheet.cell(i, col), book)
                if began is None:
                    missing_dates.append({"source": list(key), "raw_date_cell": raw["Date Event Began"], "raw_event_month": raw.get("Event Month") or raw.get("Month")})
                for field in ("Demand Loss (MW)", "Number of Customers Affected"):
                    if isinstance(raw.get(field), str) and raw[field].lower() in ("unknown", "n/a", "na", "not available"):
                        unknown_numeric_values[field] += 1
                if record:
                    check(record["event_date"] == began, "doe_event_date_changed_or_invented", list(key))
                    check(record["fields"] == raw, "doe_source_cells_changed", list(key))
                    check(record["source"].get("sha256") == digest, "doe_source_hash", list(key))
                    check(record.get("is_synthetic") is False and record.get("firm_link") is None and record.get("causal_path_gold") is None,
                          "doe_unobserved_firm_or_causal_claim", list(key))
            if boundaries:
                table_boundaries.append({"year": year, "kind": "doe417_incidents", "file": path.name, "sheet": sheet.name,
                    "source_rows": len(boundaries), "first_source_row": boundaries[0], "last_source_row": boundaries[-1]})
        doe_counts[str(year)] = raw_count
    check(set(incident_map) == raw_incident_keys, "extra_or_missing_doe_locators", {"extra": len(set(incident_map) - raw_incident_keys), "missing": len(raw_incident_keys - set(incident_map))})
    check(len(incidents) == 2006 and len(missing_dates) == 32, "doe_expected_actual_source_counts", {"incidents": len(incidents), "unknown_start_dates": len(missing_dates)})
    check(summary.get("official_incident_count") == len(incidents), "summary_incident_total", len(incidents))
    check(summary.get("incidents_with_unknown_start_date") == len(missing_dates), "summary_unknown_date_total", len(missing_dates))
    check(summary.get("firm_risk_gold_label_count") == 0 and summary.get("causal_path_gold_label_count") == 0 and summary.get("generated_record_count") == 0,
          "summary_unobserved_gold_claim", "No fabricated gold labels or generated observations")
    result = {"status": "passed" if not errors else "failed", "scope": "Independent raw spreadsheet cell-to-projection audit; no source modified",
        "eia_source_rows_by_year_and_table": source_counts, "eia_total_projected_rows": len(observations),
        "eia_projected_cells_exactly_checked": verified_cells, "eia_nonblank_rows_without_utility_id": skipped_nonblank,
        "source_supported_entities": len(entities), "entity_types": dict(Counter(record["entity_type"] for record in entities)),
        "source_supported_relations": len(relations), "relation_types": dict(Counter(record["relation_type"] for record in relations)),
        "doe_source_rows_by_year": doe_counts, "doe_incidents": len(incidents), "doe_unknown_start_dates": len(missing_dates),
        "doe_unknown_start_date_source_rows": missing_dates, "doe_unknown_numeric_cells_preserved": dict(unknown_numeric_values),
        "boundary_rows_by_year_table_sheet": table_boundaries, "raw_material_sha256": material_sha,
        "output_sha256": {path.name: sha(path) for path in sorted(OUT.iterdir()) if path.is_file()},
        "errors": dict(errors), "error_examples": dict(examples),
        "semantic_limits": ["Utilities are reported identifiers, not deduplicated parent companies.",
            "Plant and generator records include reported/planned/retired assets; they are not paper product categories.",
            "Fuel and NAICS codes are observed categories; they do not establish paper raw-material/product supply transactions.",
            "DOE incident rows are area-level reports; firm attribution, per-firm interruption labels and causal paths remain absent.",
            "Survey year is not verified publication or historical ingestion time."]}
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": result["status"], "eia_rows": len(observations), "entities": len(entities), "relations": len(relations),
                      "doe_incidents": len(incidents), "unknown_dates": len(missing_dates), "errors": dict(errors)}))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
