"""Project factual EIA/DOE table cells into traceable records, never synthetic targets."""
import collections
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

import openpyxl

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / ".runtime"))
import xlrd

ROOT = BASE / "energy_sources"
OUT = ROOT / "normalized"


def clean(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return value.strip() if isinstance(value, str) else value


def key(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and int(value) == value:
        return str(int(value))
    return str(value).strip()


def write(handle, record):
    handle.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n")


def date_value(cell, workbook):
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode).date().isoformat()
    raw = str(cell.value).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    entities, graph, counts, annual = {}, {}, collections.Counter(), {}
    warnings, incidents = [], []
    generated_record_count = 0
    def entity(identifier, kind, name, year, locator):
        record = entities.setdefault(identifier, {"entity_id": identifier, "entity_type": kind,
            "name": name, "first_observed_survey_year": year, "last_observed_survey_year": year,
            "first_source": locator, "is_synthetic": False})
        record["last_observed_survey_year"] = max(record["last_observed_survey_year"], year)
    def edge(src, dst, relation, year, locator, attributes=None):
        if src is None or dst is None:
            return
        edge_key = (src, dst, relation)
        if edge_key not in graph:
            graph[edge_key] = {"head": src, "tail": dst, "relation_type": relation,
                "first_observed_survey_year": year, "last_observed_survey_year": year,
                "first_source": locator, "is_synthetic": False}
        graph[edge_key]["last_observed_survey_year"] = max(year, graph[edge_key]["last_observed_survey_year"])
        if attributes:
            graph[edge_key].setdefault("observed_attributes", {})[str(year)] = attributes
    fields = {
        "utility": ["Utility ID", "Utility Name", "State", "Entity Type", "Owner of Plants Reported on Form?", "Operator of Plants Reported on Form?"],
        "plant": ["Utility ID", "Utility Name", "Plant Code", "Plant Name", "State", "County", "Latitude", "Longitude", "NERC Region", "Balancing Authority Code", "Primary Purpose (NAICS Code)", "Sector", "Sector Name"],
        "generator": ["Utility ID", "Utility Name", "Plant Code", "Plant Name", "State", "Generator ID", "Technology", "Prime Mover", "Status", "Nameplate Capacity (MW)", "Summer Capacity (MW)", "Winter Capacity (MW)", "Operating Month", "Operating Year", "Energy Source 1", "Energy Source 2", "Energy Source 3", "Energy Source 4", "Energy Source 5", "Energy Source 6"],
        "ownership": ["Utility ID", "Utility Name", "Plant Code", "Plant Name", "Generator ID", "Owner Name", "Ownership ID", "Percent Owned", "Status"]}
    with (OUT / "annual_observations.jsonl").open("w", encoding="utf-8") as observations:
        for year in range(2018, 2026):
            archive = ROOT / f"raw/eia860{year}.zip"
            digest = hashlib.sha256(archive.read_bytes()).hexdigest()
            year_counts = collections.Counter()
            with zipfile.ZipFile(archive) as z:
                if z.testzip() is not None:
                    raise ValueError(f"Corrupt source archive {archive}")
                for kind, prefix in [("utility", "1___Utility"), ("plant", "2___Plant"), ("generator", "3_1_Generator"), ("ownership", "4___Owner")]:
                    members = [name for name in z.namelist() if name.startswith(prefix) and name.endswith(".xlsx")]
                    if len(members) != 1:
                        raise ValueError(f"Unrecognized official table layout: {year} {prefix}")
                    member = members[0]
                    workbook = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True, data_only=True)
                    for sheet in workbook:
                        rows = sheet.iter_rows(values_only=True)
                        next(rows)
                        headers = [str(v).strip() if v is not None else "" for v in next(rows)]
                        for rownum, values in enumerate(rows, start=3):
                            raw = {name: clean(value) for name, value in zip(headers, values) if name}
                            uid = key(raw.get("Utility ID"))
                            if uid is None:
                                if any(v is not None and v != "" for v in values):
                                    warnings.append({"year": year, "member": member, "sheet": sheet.title, "row": rownum, "reason": "row has no utility ID; preserved in raw source"})
                                continue
                            locator = {"archive": f"raw/eia860{year}.zip", "archive_sha256": digest,
                                       "member": member, "sheet": sheet.title, "row": rownum}
                            projected = {name: raw[name] for name in fields[kind] if name in raw}
                            write(observations, {"survey_year": year, "table_kind": kind, "source": locator, "fields": projected})
                            year_counts[kind] += 1
                            utility = "eia:utility:" + uid
                            entity(utility, "utility_reporting_entity", raw.get("Utility Name"), year, locator)
                            plant_code = key(raw.get("Plant Code"))
                            plant = "eia:plant:" + plant_code if plant_code else None
                            gid = key(raw.get("Generator ID"))
                            generator = f"eia:generator:{plant_code}:{gid}" if plant_code and gid else None
                            if plant:
                                entity(plant, "plant", raw.get("Plant Name"), year, locator)
                                if kind == "plant":
                                    edge(plant, utility, "ReportedBy", year, locator)
                            if generator:
                                entity(generator, "generator", gid, year, locator)
                                edge(generator, plant, "LocatedAtPlant", year, locator)
                                if kind == "generator":
                                    edge(generator, utility, "ReportedBy", year, locator)
                            state = key(raw.get("State"))
                            if state:
                                region = "eia:state:" + state
                                entity(region, "region_state_code", state, year, locator)
                                edge(plant if plant else utility, region, "LocatedInState", year, locator)
                            if kind == "plant":
                                naics = key(raw.get("Primary Purpose (NAICS Code)"))
                                if naics:
                                    industry = "eia:naics:" + naics
                                    entity(industry, "industry_naics_code", naics, year, locator)
                                    edge(plant, industry, "PrimaryPurposeNAICS", year, locator)
                            if kind == "generator":
                                for field in [f"Energy Source {index}" for index in range(1, 7)]:
                                    fuel = key(raw.get(field))
                                    if fuel:
                                        target = "eia:fuel:" + fuel
                                        entity(target, "fuel_code", fuel, year, locator)
                                        edge(generator, target, "UsesEnergySource", year, locator)
                            if kind == "ownership" and generator:
                                oid = key(raw.get("Ownership ID"))
                                if oid:
                                    owner = "eia:utility:" + oid
                                    entity(owner, "utility_reporting_entity", raw.get("Owner Name"), year, locator)
                                    edge(owner, generator, "OwnsReportedGenerator", year, locator, {"source_percent_owned_cell": raw.get("Percent Owned")})
                    workbook.close()
            annual[str(year)] = dict(year_counts)
            counts.update(year_counts)
            print(json.dumps({"normalized_year": year, "records": dict(year_counts)}), flush=True)
    for path in sorted((ROOT / "raw").glob("doe417_*_Annual_Summary.xls")):
        year = int(path.name.split("_")[1])
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = xlrd.open_workbook(path, logfile=io.StringIO())
        for sheet in workbook.sheets():
            header_row = next((i for i in range(min(10, sheet.nrows)) if "Date Event Began" in sheet.row_values(i)), None)
            if header_row is None:
                continue
            headers = [str(v).strip() for v in sheet.row_values(header_row)]
            date_col = headers.index("Date Event Began")
            for i in range(header_row + 1, sheet.nrows):
                source_fields = {name: clean(value) for name, value in zip(headers, sheet.row_values(i)) if name}
                began = date_value(sheet.cell(i, date_col), workbook)
                if began is None:
                    if not ((source_fields.get("Alert Criteria") or source_fields.get("Event Type")) and
                            (source_fields.get("Area Affected") or source_fields.get("NERC Region"))):
                        continue
                    warnings.append({"source": path.name, "row": i + 1,
                                     "reason": "real incident retained with unknown/unparseable reported start date"})
                if began is not None and int(began[:4]) != year:
                    warnings.append({"source": path.name, "row": i + 1, "reason": "event date outside annual file year"})
                record = {"event_id": f"doe417:{year}:{sheet.name}:{i+1}",
                    "event_date": began, "event_kind": "reported_electric_disturbance",
                    "source_survey_year": year,
                    "date_quality": "reported_start_date" if began is not None else "unknown_reported_start_date",
                    "fields": source_fields,
                    "source": {"file": str(path.relative_to(ROOT)).replace("\\", "/"), "sha256": digest, "sheet": sheet.name, "row": i + 1},
                    "source_url": f"https://doe417.pnnl.gov/summaries/{year}_Annual_Summary.xls",
                    "firm_link": None, "causal_path_gold": None, "is_synthetic": False}
                incidents.append(record)
    for name, records in [("entities.jsonl", entities.values()), ("relations.jsonl", graph.values()), ("risk_events.jsonl", incidents)]:
        with (OUT / name).open("w", encoding="utf-8") as handle:
            for record in records:
                write(handle, record)
    types = dict(collections.Counter(r["entity_type"] for r in entities.values()))
    missing_endpoints = sum(r["head"] not in entities or r["tail"] not in entities for r in graph.values())
    assert missing_endpoints == 0
    summary = {"is_synthetic": False, "source": "EIA-860 annual 2018-2025; DOE-417 annual 2018-2023",
        "entity_count": len(entities), "entity_types": types, "unique_fact_relation_count": len(graph),
        "relation_types": dict(collections.Counter(r["relation_type"] for r in graph.values())),
        "annual_record_count": sum(counts.values()), "annual_table_rows": annual,
        "official_incident_count": len(incidents),
        "incident_counts_by_year": dict(collections.Counter(r["event_date"][:4] if r["event_date"] else "unknown" for r in incidents)),
        "incident_counts_by_source_year": dict(collections.Counter(str(r["source_survey_year"]) for r in incidents)),
        "incidents_with_unknown_start_date": sum(r["event_date"] is None for r in incidents),
        "missing_incident_years": [2024, 2025], "firm_risk_gold_label_count": 0,
        "causal_path_gold_label_count": 0, "generated_record_count": generated_record_count,
        "missing_relation_endpoints": missing_endpoints, "warnings": warnings,
        "notes": ["utility IDs identify reporting entities; not deduplicated ultimate parent firms", "plant/generator nodes are physical assets, not paper product categories", "fuel codes are not a claim to cover all raw materials", "DOE incidents disclose areas, not verified per-firm disruptions", "survey year is not publication/availability time; collected in 2026", "Percent Owned cells preserved unchanged; verify source formatting before unit conversion"]}
    (OUT / "collection_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()
